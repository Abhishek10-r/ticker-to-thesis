"""Excel valuation model with LIVE formulas.

Inputs (blue) sit on the Assumptions sheet; every downstream number is a formula, so changing a
growth rate, margin, WACC input or the scenario switch reprices the whole model. Colour code:
blue = hard-coded input/data, black = formula, green = link to another sheet, yellow fill = key assumption.
"""
from __future__ import annotations

import math
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from .analysis import METRICS
from .sec import filing_url
from .standardize import LABELS

NAVY, LIGHT, GREY = "1F2A44", "E8ECF4", "F3F4F6"
F = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
M = 1e6
FMT_M = '#,##0;(#,##0);"-"'
FMT_PS = '$#,##0.00;($#,##0.00);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_X = '0.0"x";(0.0"x");"-"'
FMT_B = '0.00'
THIN = Side(style="thin", color="C9CED8")


def _font(color=BLACK, bold=False, size=10, italic=False): return Font(name=F, color=color, bold=bold, size=size, italic=italic)


class Sheet:
    def __init__(self, wb, title, widths=None):
        self.ws = wb.create_sheet(title); self.title = title
        self.ws.sheet_view.showGridLines = False
        for col, w in (widths or {}).items(): self.ws.column_dimensions[col].width = w

    def ref(self, cell, absolute=True):
        c = "".join(ch for ch in cell if ch.isalpha()); r = "".join(ch for ch in cell if ch.isdigit())
        name = f"'{self.title}'" if " " in self.title else self.title
        return f"{name}!${c}${r}" if absolute else f"{name}!{cell}"

    def put(self, cell, value, fmt=None, color=None, bold=False, fill=None, italic=False, note=None, align=None, size=10):
        c = self.ws[cell]
        c.value = None if (isinstance(value, float) and (math.isnan(value) or math.isinf(value))) else value
        if color is None:
            color = GREEN if isinstance(value, str) and value.startswith("=") and "!" in value else \
                    (BLACK if isinstance(value, str) and value.startswith("=") else (BLUE if isinstance(value, (int, float)) else BLACK))
        c.font = _font(color, bold, size, italic)
        if fmt: c.number_format = fmt
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        if note: c.comment = Comment(note, "Ticker-to-Thesis", width=320, height=120)
        if align: c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
        return c

    def header(self, row, labels, start_col=2, fill=NAVY):
        for i, lab in enumerate(labels):
            c = self.ws.cell(row=row, column=start_col + i, value=lab)
            c.font = _font("FFFFFF", True); c.fill = PatternFill("solid", fgColor=fill)
            c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)

    def title_block(self, title, subtitle):
        self.put("B1", title, color=NAVY, bold=True, size=14)
        self.put("B2", subtitle, color="555555", italic=True)

    def section(self, cell, text):
        self.put(cell, text, color=NAVY, bold=True, size=11)


def build_workbook(rep, path) -> Path:
    v, f, a, w = rep.val, rep.fin, rep.val.assumptions, rep.val.wacc
    cfg = rep.cfg
    wb = Workbook(); wb.remove(wb.active)
    n = a.years
    yc = [L(4 + i) for i in range(n)]            # forecast columns D..H
    cover = Sheet(wb, "Cover", {"A": 2, "B": 34, "C": 22, "D": 60})
    A = Sheet(wb, "Assumptions", {"A": 2, "B": 38, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 70})
    H = Sheet(wb, "Historicals", {"A": 2, "B": 34, **{L(3 + i): 12 for i in range(10)}})
    W = Sheet(wb, "WACC", {"A": 2, "B": 30, **{L(3 + i): 13 for i in range(9)}})
    D = Sheet(wb, "DCF", {"A": 2, "B": 34, "C": 15, **{c: 13 for c in yc}, "I": 50})
    S = Sheet(wb, "Sensitivity", {"A": 2, "B": 16, **{L(3 + i): 12 for i in range(12)}})
    C = Sheet(wb, "Comps", {"A": 2, "B": 9, "C": 30, **{L(4 + i): 12 for i in range(14)}})
    V = Sheet(wb, "Valuation", {"A": 2, "B": 50, "C": 14, "D": 14, "E": 14, "F": 40})
    B = Sheet(wb, "Benchmark", {"A": 2, "B": 38, **{L(3 + i): 13 for i in range(10)}})
    P = Sheet(wb, "Sources", {"A": 2, "B": 26, "C": 8, "D": 13, "E": 16, "F": 60, "G": 9, "H": 24, "I": 12})
    for sh in [cover, A, H, W, D, S, C, V, B, P]: sh.ws.sheet_properties.tabColor = NAVY if sh in (cover, V) else "9AA5B8"

    # ================= Assumptions =================
    A.title_block(f"{f.name} ({f.ticker}): assumptions", "Blue = input. Yellow = key assumption. Edit these cells; everything else recalculates.")
    A.section("B3", "Scenario")
    A.put("B4", "Active scenario (1 = Base, 2 = Bear, 3 = Bull)"); A.put("C4", 1, fill="FFFF00", align="center")
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False); A.ws.add_data_validation(dv); dv.add("C4")
    A.section("B6", "Market & valuation inputs")
    rows = [("rf", "Risk-free rate (10-year Treasury)", w.rf, FMT_PCT, w.rf_source),
            ("erp", "Equity risk premium", w.erp, FMT_PCT, "Config assumption; update from Damodaran's latest implied ERP."),
            ("tax_m", "Marginal tax rate (debt shield, beta unlevering)", w.tax, FMT_PCT, "US federal 21% + typical state."),
            ("spread", f"Credit spread (synthetic rating {w.rating})", w.spread, FMT_PCT,
             f"Interest coverage {w.coverage:.1f}x maps to {w.rating}; spread after Damodaran's table (approximate)."),
            ("tg", "Terminal growth rate", a.terminal_growth, FMT_PCT, a.notes.get("terminal_growth", "")),
            ("exit", "Exit multiple (EV / EBITDA)", a.exit_multiple, FMT_X, a.notes.get("exit_multiple", "")),
            ("floor", "Minimum WACC spread over terminal growth", cfg["dcf"]["min_wacc_spread_over_g"], FMT_PCT, "Sanity floor."),
            ("mid", "Mid-year discounting (1 = on, 0 = off)", 1 if cfg["dcf"]["mid_year_convention"] else 0, "0", "Cash flows arrive mid-year on average.")]
    R = {}
    for i, (k, lab, val, fm, note) in enumerate(rows):
        r = 7 + i; A.put(f"B{r}", lab); A.put(f"C{r}", float(val), fm, fill="FFFF00" if k in ("erp", "tg", "exit") else None)
        A.put(f"H{r}", note, color="555555", italic=True); R[k] = A.ref(f"C{r}")
    A.section("B16", f"Market data (latest balance sheet {f.latest['asof'] or 'last 10-K'}; $m except per share)")
    L_ = f.latest
    mrows = [("price", "Share price ($)", v.price, FMT_PS, f"Latest close, {rep.run_date}."),
             ("shares", "Shares outstanding (m)", L_["shares"] / M, '#,##0.0', f"{L_['shares_source']}, {L_['shares_date']}."),
             ("cash", "Cash & equivalents", L_["cash"] / M, FMT_M, "Latest 10-Q/10-K."),
             ("sti", "Short-term investments", L_["st_investments"] / M, FMT_M, "Latest 10-Q/10-K."),
             ("debt", "Total debt (excl. operating leases)", L_["debt"] / M, FMT_M, "Long-term debt incl. current portion + short-term borrowings.")]
    for i, (k, lab, val, fm, note) in enumerate(mrows):
        r = 17 + i; A.put(f"B{r}", lab); A.put(f"C{r}", float(val), fm); A.put(f"H{r}", note, color="555555", italic=True); R[k] = A.ref(f"C{r}")
    A.put("B22", "Net debt", bold=True); A.put("C22", f"={R['debt']}-{R['cash']}-{R['sti']}", FMT_M, bold=True); R["net_debt"] = A.ref("C22")
    A.put("B23", "Market capitalisation", bold=True); A.put("C23", f"={R['price']}*{R['shares']}", FMT_M, bold=True); R["mcap"] = A.ref("C23")

    A.section("B25", "Operating drivers (base case)")
    A.header(26, ["Driver"] + [f"Year {i + 1}" for i in range(n)])
    A.put("B27", "Revenue growth"); A.put("B28", "EBIT margin")
    for i, c in enumerate([L(3 + i) for i in range(n)]):
        A.put(f"{c}27", float(a.revenue_growth[i]), FMT_PCT, fill="FFFF00"); A.put(f"{c}28", float(a.ebit_margin[i]), FMT_PCT, fill="FFFF00")
    A.put("H27", a.notes["revenue_growth"], color="555555", italic=True); A.put("H28", a.notes["ebit_margin"], color="555555", italic=True)
    for i, (k, lab, val) in enumerate([("tax", "Tax rate on EBIT", a.tax_rate), ("da", "D&A % revenue", a.da_pct),
                                        ("capex", "Capex % revenue", a.capex_pct), ("nwc", "Working capital % revenue", a.nwc_pct)]):
        r = 29 + i; A.put(f"B{r}", lab); A.put(f"C{r}", float(val), FMT_PCT); A.put(f"H{r}", a.notes.get(k if k != "tax" else "tax_rate", a.notes.get(k + "_pct", "")), color="555555", italic=True)
        R[k] = A.ref(f"C{r}")
    A.put("H30", a.notes["da_pct"], color="555555", italic=True); A.put("H31", a.notes["capex_pct"], color="555555", italic=True)
    A.put("H32", a.notes["nwc_pct"], color="555555", italic=True)

    A.section("B34", "Scenario shifts (added to base case)")
    A.header(35, ["Shift", "Base", "Bear", "Bull"])
    sc = cfg["scenarios"]
    for i, (k, lab) in enumerate([("growth", "Revenue growth (each year)"), ("margin", "EBIT margin (each year)"),
                                   ("wacc", "WACC"), ("terminal_growth", "Terminal growth")]):
        r = 36 + i; A.put(f"B{r}", lab); A.put(f"C{r}", 0.0, FMT_PCT); A.put(f"D{r}", float(sc["bear"][k]), FMT_PCT); A.put(f"E{r}", float(sc["bull"][k]), FMT_PCT)
        R["s_" + k] = f"INDEX(Assumptions!$C${r}:$E${r},Assumptions!$C$4)"

    A.section("B41", "Price target weights & rating thresholds")
    used = {k: wt for k, (val, wt) in v.pt_components.items()}
    for i, (k, lab) in enumerate([("dcf_perpetuity", "Weight: DCF (perpetuity growth)"), ("dcf_exit_multiple", "Weight: DCF (exit multiple)"),
                                   ("comps_ev_ebitda", "Weight: trading comps (EV/EBITDA median)")]):
        r = 42 + i; A.put(f"B{r}", lab); A.put(f"C{r}", float(used[k]), FMT_PCT); R["w_" + k] = A.ref(f"C{r}")
    A.put("B45", "BUY if upside above"); A.put("C45", float(cfg["rating"]["buy_above"]), FMT_PCT); R["buy"] = A.ref("C45")
    A.put("B46", "SELL if upside below"); A.put("C46", float(cfg["rating"]["sell_below"]), FMT_PCT); R["sell"] = A.ref("C46")

    # ================= Historicals =================
    ann = f.annual; years = list(ann.index); hc = [L(3 + i) for i in range(len(years))]
    H.title_block(f"{f.name}: standardised historical financials ($m)", "Source: SEC EDGAR XBRL (10-K). See Sources sheet for the tag and filing behind every value.")
    H.header(4, ["Fiscal year"] + [f"FY{y}" for y in years])
    items = ["revenue", "cogs", "gross_profit", "operating_income", "d_and_a", "ebitda", "interest_expense", "pretax_income",
             "income_tax", "net_income", "cfo", "capex", "fcf", "sbc", "dividends", "buybacks", "cash", "st_investments",
             "receivables", "inventory", "payables", "total_assets", "equity", "debt", "net_debt", "nwc"]
    hr = {}
    for i, it in enumerate(items):
        r = 5 + i; hr[it] = r; H.put(f"B{r}", LABELS.get(it, it), bold=it in ("revenue", "operating_income", "ebitda", "net_income", "fcf"))
        for j, y in enumerate(years):
            val = ann.loc[y, it]
            if it == "ebitda": H.put(f"{hc[j]}{r}", f"={hc[j]}{hr['operating_income']}+{hc[j]}{hr['d_and_a']}", FMT_M)
            elif it == "fcf": H.put(f"{hc[j]}{r}", f"={hc[j]}{hr['cfo']}-{hc[j]}{hr['capex']}", FMT_M)
            elif it == "net_debt": H.put(f"{hc[j]}{r}", f"={hc[j]}{hr['debt']}-{hc[j]}{hr['cash']}-{hc[j]}{hr['st_investments']}", FMT_M)
            elif it == "nwc": H.put(f"{hc[j]}{r}", f"={hc[j]}{hr['receivables']}+{hc[j]}{hr['inventory']}-{hc[j]}{hr['payables']}", FMT_M)
            else: H.put(f"{hc[j]}{r}", None if np.isnan(val) else float(val) / M, FMT_M)
    r0 = 5 + len(items) + 1
    H.section(f"B{r0}", "Ratios (formulas)")
    ratio_rows = [("Revenue growth", lambda c, p: f"=IFERROR({c}{hr['revenue']}/{p}{hr['revenue']}-1,\"\")" if p else None),
                  ("Gross margin", lambda c, p: f"=IFERROR({c}{hr['gross_profit']}/{c}{hr['revenue']},\"\")"),
                  ("EBIT margin", lambda c, p: f"=IFERROR({c}{hr['operating_income']}/{c}{hr['revenue']},\"\")"),
                  ("EBITDA margin", lambda c, p: f"=IFERROR({c}{hr['ebitda']}/{c}{hr['revenue']},\"\")"),
                  ("Net margin", lambda c, p: f"=IFERROR({c}{hr['net_income']}/{c}{hr['revenue']},\"\")"),
                  ("FCF margin", lambda c, p: f"=IFERROR({c}{hr['fcf']}/{c}{hr['revenue']},\"\")"),
                  ("Effective tax rate", lambda c, p: f"=IFERROR({c}{hr['income_tax']}/{c}{hr['pretax_income']},\"\")"),
                  ("D&A % revenue", lambda c, p: f"=IFERROR({c}{hr['d_and_a']}/{c}{hr['revenue']},\"\")"),
                  ("Capex % revenue", lambda c, p: f"=IFERROR({c}{hr['capex']}/{c}{hr['revenue']},\"\")"),
                  ("Working capital % revenue", lambda c, p: f"=IFERROR({c}{hr['nwc']}/{c}{hr['revenue']},\"\")"),
                  ("Net debt / EBITDA", lambda c, p: f"=IFERROR({c}{hr['net_debt']}/{c}{hr['ebitda']},\"\")")]
    for i, (lab, fn) in enumerate(ratio_rows):
        r = r0 + 1 + i; H.put(f"B{r}", lab)
        for j, c in enumerate(hc):
            fml = fn(c, hc[j - 1] if j else None)
            if fml: H.put(f"{c}{r}", fml, FMT_X if "EBITDA" in lab and "/" in lab else FMT_PCT)
    H.ws.freeze_panes = "C5"
    last_rev = f"Historicals!${hc[-1]}${hr['revenue']}"

    # ================= WACC =================
    W.title_block(f"{f.ticker}: cost of capital", "Bottom-up beta: peers' regression betas, unlevered, median, relevered at the target's capital structure.")
    W.header(4, ["Company", "Raw beta (2y weekly)", "R²", "Adjusted beta", "Market cap ($m)", "Debt ($m)", "Debt / equity", "Unlevered beta", "Role"])
    bt = w.beta_table.copy(); order = [rep.target] + [t for t in bt.index if t != rep.target]; bt = bt.loc[[t for t in order if t in bt.index]]
    blume = cfg["beta"]["blume_adjust"]
    caps = v.comps["market_cap"]
    for i, (tk, row) in enumerate(bt.iterrows()):
        r = 5 + i
        W.put(f"B{r}", tk, bold=tk == rep.target); W.put(f"C{r}", float(row["raw_beta"]), FMT_B); W.put(f"D{r}", float(row["r2"]), FMT_B)
        W.put(f"E{r}", (f"=IF(C{r}=\"\",\"\",0.67*C{r}+0.33)" if blume else f"=IF(C{r}=\"\",\"\",C{r})"), FMT_B)
        W.put(f"F{r}", float(caps[tk]) / M, FMT_M); W.put(f"G{r}", float(rep.fins[tk].latest["debt"]) / M, FMT_M)
        W.put(f"H{r}", f"=IFERROR(G{r}/F{r},0)", FMT_PCT); W.put(f"I{r}", f"=IF(E{r}=\"\",\"\",E{r}/(1+(1-{R['tax_m']})*H{r}))", FMT_B)
        W.put(f"J{r}", "Target" if tk == rep.target else "Peer")
    last = 4 + len(bt); p0 = 6
    rr = last + 2
    lines = [("Median unlevered beta (peers)", f"=MEDIAN(I{p0}:I{last})", FMT_B),
             ("Target debt / equity (market)", f"={R['debt']}/{R['mcap']}", FMT_PCT),
             ("Relevered beta", f"=C{rr}*(1+(1-{R['tax_m']})*C{rr + 1})", FMT_B),
             ("Risk-free rate", f"={R['rf']}", FMT_PCT), ("Equity risk premium", f"={R['erp']}", FMT_PCT),
             ("Cost of equity", f"=C{rr + 3}+C{rr + 2}*C{rr + 4}", FMT_PCT),
             ("Pre-tax cost of debt (rf + spread)", f"={R['rf']}+{R['spread']}", FMT_PCT),
             ("After-tax cost of debt", f"=C{rr + 6}*(1-{R['tax_m']})", FMT_PCT),
             ("Weight of equity", f"={R['mcap']}/({R['mcap']}+{R['debt']})", FMT_PCT),
             ("Weight of debt", f"=1-C{rr + 8}", FMT_PCT),
             ("WACC (calculated)", f"=C{rr + 8}*C{rr + 5}+C{rr + 9}*C{rr + 7}", FMT_PCT),
             ("WACC (after floor: terminal growth + min spread)", f"=MAX(C{rr + 10},{R['tg']}+{R['floor']})", FMT_PCT),
             ("WACC used (incl. scenario shift)", f"=C{rr + 11}+{R['s_wacc']}", FMT_PCT)]
    for i, (lab, fml, fm) in enumerate(lines):
        r = rr + i; W.put(f"B{r}", lab, bold=i in (5, 12)); W.put(f"C{r}", fml, fm, bold=i in (5, 12))
    W.put(f"D{rr + 2}", f"Target's own adjusted beta: {w.beta_own_adjusted:.2f} (shown for reference)", color="555555", italic=True)
    R["wacc"] = W.ref(f"C{rr + 12}")

    # ================= DCF =================
    D.title_block(f"{f.ticker}: discounted cash flow ($m)", "Unlevered free cash flow = EBIT × (1 − tax) + D&A − capex − Δ working capital. Mid-year discounting.")
    D.header(4, ["", f"FY{years[-1]}A"] + [f"Year {i + 1}" for i in range(n)])
    D.put("B5", "Year number"); [D.put(f"{c}5", i + 1, "0") for i, c in enumerate(yc)]
    D.put("B6", "Revenue", bold=True); D.put("C6", f"={last_rev}", FMT_M, bold=True)
    labels = {7: "Revenue growth", 8: "EBIT margin", 9: "EBIT", 10: "Taxes on EBIT", 11: "NOPAT", 12: "D&A", 13: "Capex",
              14: "Working capital", 15: "Δ Working capital", 16: "Unlevered free cash flow", 17: "EBITDA",
              18: "Discount period (years)", 19: "Discount factor", 20: "PV of UFCF"}
    for r, lab in labels.items(): D.put(f"B{r}", lab, bold=r in (16, 20))
    D.put("C14", f"=C6*{R['nwc']}", FMT_M)
    for i, c in enumerate(yc):
        p = "C" if i == 0 else yc[i - 1]; ac = L(3 + i)
        D.put(f"{c}6", f"={p}6*(1+{c}7)", FMT_M, bold=True)
        D.put(f"{c}7", f"=Assumptions!{ac}27+{R['s_growth']}", FMT_PCT)
        D.put(f"{c}8", f"=Assumptions!{ac}28+{R['s_margin']}", FMT_PCT)
        D.put(f"{c}9", f"={c}6*{c}8", FMT_M); D.put(f"{c}10", f"={c}9*{R['tax']}", FMT_M)
        D.put(f"{c}11", f"={c}9-{c}10", FMT_M); D.put(f"{c}12", f"={c}6*{R['da']}", FMT_M)
        D.put(f"{c}13", f"={c}6*{R['capex']}", FMT_M); D.put(f"{c}14", f"={c}6*{R['nwc']}", FMT_M)
        D.put(f"{c}15", f"={c}14-{p}14", FMT_M); D.put(f"{c}16", f"={c}11+{c}12-{c}13-{c}15", FMT_M, bold=True)
        D.put(f"{c}17", f"={c}9+{c}12", FMT_M); D.put(f"{c}18", f"={c}5-0.5*{R['mid']}", "0.0")
        D.put(f"{c}19", f"=1/(1+$C$23)^{c}18", "0.000"); D.put(f"{c}20", f"={c}16*{c}19", FMT_M, bold=True)
    N = yc[-1]
    D.section("B22", "Perpetuity growth method")
    blk = [(23, "WACC", f"={R['wacc']}", FMT_PCT), (24, "Terminal growth", f"={R['tg']}+{R['s_terminal_growth']}", FMT_PCT),
           (25, "Sum of PV of UFCF", f"=SUM(D20:{N}20)", FMT_M),
           (26, "Terminal value", f"={N}16*(1+C24)/(C23-C24)", FMT_M), (27, "PV of terminal value", f"=C26/(1+C23)^{N}5", FMT_M),
           (28, "Enterprise value", "=C25+C27", FMT_M), (29, "Less: net debt", f"={R['net_debt']}", FMT_M),
           (30, "Equity value", "=C28-C29", FMT_M), (31, "Shares outstanding (m)", f"={R['shares']}", '#,##0.0'),
           (32, "Value per share ($)", "=C30/C31", FMT_PS), (33, "Terminal value % of EV", "=C27/C28", FMT_PCT),
           (34, "Implied exit multiple (EV/EBITDA)", f"=C26/{N}17", FMT_X)]
    for r, lab, fml, fm in blk: D.put(f"B{r}", lab, bold=r in (28, 32)); D.put(f"C{r}", fml, fm, bold=r in (28, 32))
    D.section("B36", "Exit multiple method")
    blk2 = [(37, "Exit multiple (EV/EBITDA)", f"={R['exit']}", FMT_X), (38, "Terminal value", f"={N}17*C37", FMT_M),
            (39, "PV of terminal value", f"=C38/(1+C23)^{N}5", FMT_M), (40, "Enterprise value", "=C25+C39", FMT_M),
            (41, "Equity value", "=C40-C29", FMT_M), (42, "Value per share ($)", "=C41/C31", FMT_PS),
            (43, "Implied perpetual growth", f"=(C38*C23-{N}16)/(C38+{N}16)", FMT_PCT)]
    for r, lab, fml, fm in blk2: D.put(f"B{r}", lab, bold=r in (40, 42)); D.put(f"C{r}", fml, fm, bold=r in (40, 42))
    D.put("I26", "Gordon growth: UFCF in year N × (1 + g) / (WACC − g), discounted from the end of year N.", color="555555", italic=True)
    D.ws.freeze_panes = "C5"

    # ================= Sensitivity =================
    S.title_block(f"{f.ticker}: value per share sensitivity ($)", "Live formulas: each cell re-runs the DCF for that WACC and terminal assumption.")
    S.put("B4", "WACC step"); S.put("C4", 0.005, FMT_PCT); S.put("D4", "g step"); S.put("E4", 0.0025, FMT_PCT)
    S.put("F4", "Multiple step"); S.put("G4", 1.0, FMT_X)
    pv_expr = lambda wc: f"SUMPRODUCT(DCF!$D$16:${N}$16/(1+{wc})^DCF!$D$18:${N}$18)"
    S.section("B6", "Perpetuity growth method: WACC (rows) × terminal growth (columns)")
    S.put("B7", "WACC \\ g", bold=True)
    for k in range(5):
        S.put(f"{L(3 + k)}7", f"=DCF!$C$24+({k - 2})*$E$4", FMT_PCT, bold=True, fill=LIGHT)
        S.put(f"B{8 + k}", f"=DCF!$C$23+({k - 2})*$C$4", FMT_PCT, bold=True, fill=LIGHT)
    for i in range(5):
        for k in range(5):
            r, c = 8 + i, L(3 + k)
            S.put(f"{c}{r}", f"=IF($B{r}>{c}$7,({pv_expr(f'$B{r}')}+DCF!${N}$16*(1+{c}$7)/($B{r}-{c}$7)/(1+$B{r})^DCF!${N}$5-DCF!$C$29)/DCF!$C$31,\"n/a\")", FMT_PS)
    S.section("B15", "Exit multiple method: WACC (rows) × EV/EBITDA exit multiple (columns)")
    S.put("B16", "WACC \\ multiple", bold=True)
    for k in range(5):
        S.put(f"{L(3 + k)}16", f"=DCF!$C$37+({k - 2})*$G$4", FMT_X, bold=True, fill=LIGHT)
        S.put(f"B{17 + k}", f"=DCF!$C$23+({k - 2})*$C$4", FMT_PCT, bold=True, fill=LIGHT)
    for i in range(5):
        for k in range(5):
            r, c = 17 + i, L(3 + k)
            S.put(f"{c}{r}", f"=({pv_expr(f'$B{r}')}+DCF!${N}$17*{c}$16/(1+$B{r})^DCF!${N}$5-DCF!$C$29)/DCF!$C$31", FMT_PS)
    for rng in ("C8:G12", "C17:G21"):
        S.ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="F4CCCC", mid_type="percentile",
                                                            mid_value=50, mid_color="FFFFFF", end_type="max", end_color="C9E7D2"))
    S.put("B23", "Current share price", bold=True); S.put("C23", f"={R['price']}", FMT_PS)

    # ================= Comps =================
    C.title_block("Trading comparables (latest fiscal year; $m except per share)",
                  "Multiples on the latest reported fiscal year. Negative or extreme multiples are blanked (not meaningful).")
    C.header(4, ["Ticker", "Company", "FY", "Price ($)", "Shares (m)", "Market cap", "Debt", "Cash + ST inv.", "Enterprise value",
                 "Revenue", "EBITDA", "Net income", "EV / Revenue", "EV / EBITDA", "P / E", "EBIT margin"])
    comps = v.comps; order = [t for t in comps.index if t != rep.target] + [rep.target]
    mx, mp = cfg["comps"]["max_ev_ebitda"], cfg["comps"]["max_pe"]
    for i, tk in enumerate(order):
        r = 5 + i; row = comps.loc[tk]; ff = rep.fins[tk]; tgt = tk == rep.target
        fill = "FFF2CC" if tgt else None
        vals = [(tk, None), (ff.name, None), (int(row["fiscal_year"]), "0"), (float(row["price"]), FMT_PS), (ff.latest["shares"] / M, '#,##0.0'),
                (f"=E{r}*F{r}", FMT_M), (ff.latest["debt"] / M, FMT_M), ((ff.latest["cash"] + ff.latest["st_investments"]) / M, FMT_M),
                (f"=G{r}+H{r}-I{r}", FMT_M), (float(row["revenue"]) / M, FMT_M), (float(row["ebitda"]) / M, FMT_M),
                (float(row["net_income"]) / M, FMT_M), (f"=IFERROR(J{r}/K{r},\"\")", FMT_X),
                (f"=IF(AND(L{r}>0,J{r}/L{r}<={mx},J{r}>0),J{r}/L{r},\"\")", FMT_X),
                (f"=IF(AND(M{r}>0,G{r}/M{r}<={mp}),G{r}/M{r},\"\")", FMT_X), (float(row["ebit_margin"]), FMT_PCT)]
        for j, (val, fm) in enumerate(vals):
            C.put(f"{L(2 + j)}{r}", val, fm, fill=fill, bold=tgt)
    lp = 4 + len(order) - 1; tr = lp + 1           # last peer row, target row
    sr = tr + 2
    for k, (lab, fn) in enumerate([("Peer 25th percentile", "QUARTILE({},1)"), ("Peer median", "MEDIAN({})"),
                                    ("Peer 75th percentile", "QUARTILE({},3)")]):
        r = sr + k; C.put(f"C{r}", lab, bold=True)
        for col in ("N", "O", "P", "Q"):
            C.put(f"{col}{r}", f"=IFERROR({fn.format(f'{col}5:{col}{lp}')},\"\")", FMT_PCT if col == "Q" else FMT_X, bold=True, fill=GREY)
    C.section(f"B{sr + 4}", f"Implied value per share for {rep.target} ($)")
    C.header(sr + 5, ["", "Method", "Low (Q1)", "Median", "High (Q3)"])
    imp = {}
    for k, (lab, col, metric, is_pe) in enumerate([("EV / EBITDA", "O", "L", False), ("P / E", "P", "M", True), ("EV / Revenue", "N", "K", False)]):
        r = sr + 6 + k; C.put(f"C{r}", lab, bold=True)
        for j, srow in enumerate([sr, sr + 1, sr + 2]):
            cell = f"{L(4 + j)}{r}"
            fml = (f"=IFERROR({col}{srow}*{metric}{tr}/F{tr},\"\")" if is_pe
                   else f"=IFERROR(({col}{srow}*{metric}{tr}-({R['net_debt']}))/F{tr},\"\")")
            C.put(cell, fml, FMT_PS)
        imp[lab] = r
    C.ws.freeze_panes = "D5"

    # ================= Valuation summary =================
    V.title_block(f"{f.name} ({f.ticker}): valuation summary", "Football field, blended price target and rating, all linked to the model.")
    V.header(4, ["Method", "Low ($)", "High ($)", "Midpoint ($)", "Basis"])
    s52 = v.week52
    ff_rows = [("52-week trading range (closing prices)", float(s52[0]), float(s52[1]), "Market data"),
               ("Trading comps: EV / EBITDA (peer Q1–Q3)", f"=Comps!$D${imp['EV / EBITDA']}", f"=Comps!$F${imp['EV / EBITDA']}", "Comps sheet"),
               ("Trading comps: P / E (peer Q1–Q3)", f"=Comps!$D${imp['P / E']}", f"=Comps!$F${imp['P / E']}", "Comps sheet"),
               ("DCF: perpetuity growth (WACC ±0.5pp, g ±0.25pp)", "=MIN(Sensitivity!$D$9:$F$11)", "=MAX(Sensitivity!$D$9:$F$11)", "Sensitivity sheet"),
               ("DCF: exit multiple (WACC ±0.5pp, multiple ±1x)", "=MIN(Sensitivity!$D$18:$F$20)", "=MAX(Sensitivity!$D$18:$F$20)", "Sensitivity sheet"),
               ("Scenarios: bear to bull (DCF snapshot)", float(v.scenarios["Bear"].price_perpetuity), float(v.scenarios["Bull"].price_perpetuity),
                "Snapshot; switch the scenario on Assumptions to recompute live")]
    for i, (lab, lo, hi, basis) in enumerate(ff_rows):
        r = 5 + i; V.put(f"B{r}", lab); V.put(f"C{r}", lo, FMT_PS); V.put(f"D{r}", hi, FMT_PS)
        V.put(f"E{r}", f"=IFERROR((C{r}+D{r})/2,\"\")", FMT_PS); V.put(f"F{r}", basis, color="555555", italic=True)
    V.section("B12", "Blended price target")
    V.header(13, ["Component", "Value ($)", "Weight", "Contribution"])
    comp_rows = [("DCF: perpetuity growth", "=DCF!$C$32", R["w_dcf_perpetuity"]),
                 ("DCF: exit multiple", "=DCF!$C$42", R["w_dcf_exit_multiple"]),
                 ("Trading comps: EV / EBITDA median", f"=Comps!$E${imp['EV / EBITDA']}", R["w_comps_ev_ebitda"])]
    for i, (lab, val, wt) in enumerate(comp_rows):
        r = 14 + i; V.put(f"B{r}", lab); V.put(f"C{r}", val, FMT_PS); V.put(f"D{r}", f"={wt}", FMT_PCT)
        V.put(f"E{r}", f"=IFERROR(C{r}*D{r},0)", FMT_PS)
    V.put("B17", "Price target", bold=True); V.put("C17", "=SUM(E14:E16)/SUMPRODUCT(--ISNUMBER(C14:C16),D14:D16)", FMT_PS, bold=True, fill="FFF2CC")
    V.put("B18", "Current share price"); V.put("C18", f"={R['price']}", FMT_PS)
    V.put("B19", "Implied upside / (downside)", bold=True); V.put("C19", "=C17/C18-1", FMT_PCT, bold=True)
    V.put("B20", "Rating", bold=True)
    V.put("C20", f"=IF(C19>{R['buy']},\"BUY\",IF(C19<{R['sell']},\"SELL\",\"HOLD\"))", bold=True, align="center", fill="FFF2CC")
    V.section("B22", "Scenario snapshot (DCF perpetuity method)")
    V.header(23, ["Scenario", "Value per share ($)", "vs price"])
    for i, k in enumerate(["Bear", "Base", "Bull"]):
        r = 24 + i; V.put(f"B{r}", k); V.put(f"C{r}", float(v.scenarios[k].price_perpetuity), FMT_PS); V.put(f"D{r}", f"=C{r}/$C$18-1", FMT_PCT)
    # football field chart: invisible "low" series + visible range series
    V.put("H4", "Chart helper", color="999999", italic=True); V.put("H5", "Low"); V.put("I5", "Range")
    for i in range(6):
        r = 6 + i; V.put(f"G{r}", f"=B{5 + i}", color="999999"); V.put(f"H{r}", f"=C{5 + i}", FMT_PS, color="999999")
        V.put(f"I{r}", f"=D{5 + i}-C{5 + i}", FMT_PS, color="999999")
    ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100
    ch.title = "Valuation football field ($ per share)"; ch.height, ch.width = 8, 18
    ch.add_data(Reference(V.ws, min_col=8, max_col=9, min_row=5, max_row=11), titles_from_data=True)
    ch.set_categories(Reference(V.ws, min_col=7, min_row=6, max_row=11))
    ch.series[0].graphicalProperties.noFill = True; ch.series[0].graphicalProperties.line.noFill = True
    ch.series[1].graphicalProperties.solidFill = "2A78D6"; ch.legend = None
    ch.y_axis.majorGridlines = None; ch.x_axis.scaling.orientation = "maxMin"
    V.ws.add_chart(ch, "B29")

    # ================= Benchmark =================
    B.title_block(f"{f.ticker} vs peers: benchmarking", f"Peers: {', '.join(rep.peers)}. Latest fiscal year for each company.")
    B.header(4, ["Metric", rep.target, "Peer median", "Peer Q1", "Peer Q3", "Percentile (100 = best)", "Verdict"])
    for i, (m, row) in enumerate(rep.bench.scorecard.iterrows()):
        r = 5 + i; fm = {"pct": FMT_PCT, "x": FMT_X, "days": "0"}[row["format"]]
        B.put(f"B{r}", row["label"])
        for j, k in enumerate(["target", "peer_median", "peer_q1", "peer_q3"]): B.put(f"{L(3 + j)}{r}", float(row[k]), fm)
        B.put(f"G{r}", float(row["percentile"]), "0"); B.put(f"H{r}", row["verdict"], align="center",
              fill={"Better": "D9EAD3", "Worse": "F4CCCC", "In line": GREY}.get(row["verdict"]))
    r = 6 + len(rep.bench.scorecard)
    for m, tbl in rep.bench.trends.items():
        r += 1; B.section(f"B{r}", f"Trend: {METRICS[m][0]}"); r += 1
        B.header(r, ["Fiscal year"] + [f"FY{int(y)}" for y in tbl.index]); fm = FMT_X if METRICS[m][1] == "x" else FMT_PCT
        for lab, col in [(rep.target, rep.target), ("Peer median", "peer_median"), ("Peer Q1", "peer_q1"), ("Peer Q3", "peer_q3")]:
            r += 1; B.put(f"B{r}", lab, bold=col == rep.target)
            for j, y in enumerate(tbl.index): B.put(f"{L(3 + j)}{r}", float(tbl.loc[y, col]), fm)
        r += 1

    # ================= Sources =================
    P.title_block("Data provenance", "Every historical value: the XBRL tag it came from and the SEC filing (click to open).")
    P.header(4, ["Line item", "FY", "Period end", "Value ($m)", "XBRL tag", "Form", "Accession", "Filed"])
    prov = f.provenance.sort_values(["item", "fy"])
    for i, row in enumerate(prov.itertuples()):
        r = 5 + i
        P.put(f"B{r}", LABELS.get(row.item, row.item)); P.put(f"C{r}", int(row.fy), "0"); P.put(f"D{r}", str(row.period_end) if row.period_end else "derived")
        P.put(f"E{r}", float(row.value) / (M if "shares" not in row.item and "eps" not in row.item else 1), '#,##0.00')
        P.put(f"F{r}", row.tag); P.put(f"G{r}", row.form)
        c = P.put(f"H{r}", row.accn, color="0563C1"); c.hyperlink = filing_url(f.cik, row.accn)
        P.put(f"I{r}", row.filed)
    P.ws.freeze_panes = "B5"

    # ================= Cover =================
    cover.put("B2", "TICKER-TO-THESIS · EQUITY RESEARCH MODEL", color="777777", bold=True)
    cover.put("B4", f"{f.name} ({f.ticker})", color=NAVY, bold=True, size=18)
    cover.put("B5", f"{f.sic_description} · fiscal year ends {f.fiscal_year_end[:2]}/{f.fiscal_year_end[2:]} · model date {rep.run_date}"
              + (f" · {rep.analyst}" if rep.analyst else ""), color="555555", italic=True)
    kv = [("Rating", "=Valuation!$C$20", None), ("Price target ($)", "=Valuation!$C$17", FMT_PS), ("Current price ($)", f"={R['price']}", FMT_PS),
          ("Upside / (downside)", "=Valuation!$C$19", FMT_PCT), ("Market cap ($m)", f"={R['mcap']}", FMT_M),
          ("Enterprise value ($m)", f"={R['mcap']}+{R['net_debt']}", FMT_M), ("WACC", f"={R['wacc']}", FMT_PCT),
          ("DCF value per share, perpetuity ($)", "=DCF!$C$32", FMT_PS), ("DCF value per share, exit multiple ($)", "=DCF!$C$42", FMT_PS)]
    for i, (lab, fml, fm) in enumerate(kv):
        r = 7 + i; cover.put(f"B{r}", lab, bold=True); cover.put(f"C{r}", fml, fm, bold=True, align="right")
        for col in "BC": cover.ws[f"{col}{r}"].border = Border(bottom=THIN)
    cover.section("B18", "How to use this model")
    guide = ["Assumptions: change blue/yellow cells (growth, margins, ERP, terminal growth, exit multiple) or switch the scenario.",
             "DCF and Sensitivity recalculate live; Valuation shows the football field, blended price target and rating.",
             "Comps: peer multiples and implied values. Benchmark: target vs peer distribution and trends.",
             "Sources: the XBRL tag and SEC filing behind every historical number.",
             "Colour code: blue = input, black = formula, green = link to another sheet, yellow fill = key assumption."]
    for i, g in enumerate(guide): cover.put(f"B{19 + i}", f"{i + 1}. {g}")
    if rep.data_note: cover.put("B25", rep.data_note, color="C00000", bold=True)
    cover.put("B27", "Educational project generated from public SEC filings and market data. Not investment advice.", color="777777", italic=True)

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 90
    wb.calculation.fullCalcOnLoad = True          # Excel computes every formula when the file is opened
    path = Path(path); wb.save(path)
    return path

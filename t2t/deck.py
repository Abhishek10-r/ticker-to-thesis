"""Pitch deck (PowerPoint) with native, editable charts and tables.

Built with python-pptx so it runs anywhere Python runs (Colab, Streamlit Cloud). Every number on every
slide comes from the same Report object as the Excel model.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

from .analysis import METRICS, fmt

NAVY, INK, MUTED, LINE, PANEL = "14213D", "1F2933", "5F6B7A", "D5DAE1", "F2F4F7"
ACCENT, PEER, GOOD, BAD, NEUTRAL = "E07A1F", "2F5C8F", "1B8A5A", "C0392B", "8A94A3"
HEAD, BODY = "Cambria", "Calibri"
W, H = 13.333, 7.5
RATING_COL = {"BUY": GOOD, "HOLD": "B7791F", "SELL": BAD}


def rgb(h): return RGBColor.from_string(h)


class Deck:
    def __init__(self, rep):
        self.rep = rep
        self.prs = Presentation(); self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]; self.n = 0

    # ---------------- primitives ----------------
    def slide(self, dark=False):
        s = self.prs.slides.add_slide(self.blank); self.n += 1
        if dark:
            bg = s.background.fill; bg.solid(); bg.fore_color.rgb = rgb(NAVY)
        return s

    def text(self, s, x, y, w, h, runs, size=14, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, name=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: tb.name = name
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0); tf.margin_top = tf.margin_bottom = Inches(0.02)
        paras = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs if isinstance(runs, list) else [(runs, {})]]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
            for item in para:
                t, o = (item, {}) if isinstance(item, str) else item
                r = p.add_run(); r.text = t
                r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", bold); r.font.italic = o.get("italic", False)
                r.font.color.rgb = rgb(o.get("color", color)); r.font.name = o.get("font", font)
        return tb

    def bullets(self, s, x, y, w, h, items, size=14, color=INK, gap=8, marker_color=ACCENT):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0)
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(gap)
            head, body = (it if isinstance(it, tuple) else (None, it))
            r0 = p.add_run(); r0.text = "■  "; r0.font.size = Pt(size - 4); r0.font.color.rgb = rgb(marker_color); r0.font.name = BODY
            if head:
                r1 = p.add_run(); r1.text = head + " "; r1.font.bold = True; r1.font.size = Pt(size); r1.font.color.rgb = rgb(color); r1.font.name = BODY
            r2 = p.add_run(); r2.text = body; r2.font.size = Pt(size); r2.font.color.rgb = rgb(color); r2.font.name = BODY
        return tb

    def rect(self, s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, name=None):
        r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
        if line: r.line.color.rgb = rgb(line); r.line.width = Pt(0.75)
        else: r.line.fill.background()
        r.shadow.inherit = False
        if name: r.name = name
        return r

    def vline(self, s, x, y1, y2, color, width=2, dash=False):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y1), Inches(x), Inches(y2))
        c.line.color.rgb = rgb(color); c.line.width = Pt(width)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return c

    def title(self, s, headline, kicker=None):
        if kicker: self.text(s, 0.6, 0.35, 12, 0.3, kicker.upper(), size=11, color=ACCENT, bold=True)
        self.text(s, 0.6, 0.62, 12.1, 0.95, headline, size=26, color=NAVY, bold=True, font=HEAD, name="Title")

    def footer(self, s, source="Source: SEC EDGAR (10-K/10-Q XBRL), market data; Ticker-to-Thesis analysis."):
        rep = self.rep
        self.text(s, 0.6, 7.0, 9.5, 0.3, source, size=10, color=MUTED)
        self.text(s, 10.2, 7.0, 2.55, 0.3, f"{rep.target} · {rep.run_date} · {self.n}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)

    def table(self, s, x, y, w, rows, col_w=None, size=12, header_fill=NAVY, row_h=0.36, highlight_rows=(), first_bold=True,
              aligns=None, fills=None):
        nr, nc = len(rows), len(rows[0])
        gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
        if col_w:
            for j, cw in enumerate(col_w): gt.columns[j].width = Inches(cw)
        for i, row in enumerate(rows):
            gt.rows[i].height = Inches(row_h)
            for j, val in enumerate(row):
                c = gt.cell(i, j); c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]
                p.alignment = (aligns[j] if aligns else (PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT))
                r = p.add_run(); r.text = str(val); r.font.name = BODY; r.font.size = Pt(size)
                c.fill.solid()
                if i == 0:
                    c.fill.fore_color.rgb = rgb(header_fill); r.font.color.rgb = rgb("FFFFFF"); r.font.bold = True
                else:
                    f = (fills or {}).get((i, j))
                    c.fill.fore_color.rgb = rgb(f or ("FFF4E8" if i in highlight_rows else ("FFFFFF" if i % 2 else PANEL)))
                    r.font.color.rgb = rgb(INK); r.font.bold = (j == 0 and first_bold) or i in highlight_rows
        return gt

    def chart(self, s, kind, x, y, w, h, categories, series, colors, number_format='0%', legend=True, labels=False,
              title=None, gap=60, label_fmt=None, line_width=2.25):
        cd = CategoryChartData(); cd.categories = categories
        for name, vals in series: cd.add_series(name, [None if (v is None or (isinstance(v, float) and np.isnan(v))) else float(v) for v in vals])
        gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd); ch = gf.chart
        ch.font.size = Pt(12); ch.font.name = BODY; ch.font.color.rgb = rgb(MUTED)
        ch.has_title = bool(title)
        if title:
            ch.chart_title.text_frame.text = title
            tr = ch.chart_title.text_frame.paragraphs[0].runs[0]; tr.font.size = Pt(13); tr.font.bold = True; tr.font.color.rgb = rgb(INK)
        ch.has_legend = legend and len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(12)
        va = ch.value_axis; va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb("E3E7ED"); va.format.line.fill.background()
        va.tick_labels.number_format = number_format; va.tick_labels.number_format_is_linked = False
        ca = ch.category_axis; ca.format.line.color.rgb = rgb(LINE); ca.has_major_gridlines = False
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW          # keep labels below the plot when values go negative
        for i, ser in enumerate(ch.plots[0].series):
            col = colors[i % len(colors)]
            if kind in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
                ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(line_width); ser.smooth = False
                ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col)
                ser.marker.format.line.color.rgb = rgb(col); ser.marker.size = 7
            else:
                ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
        if kind in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED):
            ch.plots[0].gap_width = gap
        if labels:
            pl = ch.plots[0]; pl.has_data_labels = True; dl = pl.data_labels
            dl.number_format = label_fmt or number_format; dl.number_format_is_linked = False
            dl.font.size = Pt(11); dl.font.color.rgb = rgb(INK)
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        return ch


def _pct(x): return fmt(x, "pct")


def _money(x): return "n/a" if x is None or np.isnan(x) else f"${x:,.2f}"


def _bn(x):
    if x is None or np.isnan(x): return "n/a"
    s = f"${abs(x) / 1e9:,.1f}bn" if abs(x) >= 1e9 else f"${abs(x) / 1e6:,.0f}m"
    return f"({s})" if x < 0 else s


def build_deck(rep, path) -> Path:
    d = Deck(rep); v, b, f, a = rep.val, rep.bench, rep.fin, rep.val.assumptions
    ins = rep.insights(); av = rep.analyst_view or {}
    t = rep.target; peers = rep.peers; sc = b.scorecard
    rcol = RATING_COL[v.rating]

    # 1 ---------------- title
    s = d.slide(dark=True)
    d.text(s, 0.8, 0.9, 11, 0.4, "EQUITY RESEARCH · INITIATION OF COVERAGE", size=14, color=ACCENT, bold=True)
    d.text(s, 0.8, 1.5, 11.5, 1.2, f.name, size=44, color="FFFFFF", bold=True, font=HEAD)
    d.text(s, 0.8, 2.65, 11.5, 0.5, f"{t} · {f.sic_description.title()} · benchmarked against {', '.join(peers)}", size=16, color="C9D1DE")
    cards = [("RATING", v.rating, rcol), ("PRICE TARGET", _money(v.price_target), NAVY), ("CURRENT PRICE", _money(v.price), NAVY),
             ("IMPLIED RETURN", f"{v.upside:+.1%}", GOOD if v.upside > 0 else BAD)]
    for i, (lab, val, col) in enumerate(cards):
        x = 0.8 + i * 2.95
        d.rect(s, x, 3.75, 2.7, 1.55, "FFFFFF", shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
        d.text(s, x + 0.25, 3.92, 2.3, 0.3, lab, size=11, color=MUTED, bold=True)
        d.text(s, x + 0.25, 4.3, 2.3, 0.8, val, size=30, color=col, bold=True, font=HEAD)
    d.text(s, 0.8, 5.9, 11.5, 0.4, f"{rep.analyst + ' · ' if rep.analyst else ''}{rep.run_date}", size=14, color="FFFFFF")
    d.text(s, 0.8, 6.3, 11.5, 0.4, rep.data_note or "Built from SEC EDGAR filings and public market data with Ticker-to-Thesis. Not investment advice.",
           size=11, color="C0392B" if rep.data_note else "8FA0B8", bold=bool(rep.data_note))

    # 2 ---------------- investment summary
    s = d.slide()
    up_word = "upside" if v.upside >= 0 else "downside"
    d.title(s, f"{v.rating}: ${v.price_target:,.2f} price target implies {abs(v.upside):.0%} {up_word}", "Investment summary")
    for i, (lab, val, col) in enumerate([("Rating", v.rating, rcol), ("Price target", _money(v.price_target), NAVY),
                                         ("Implied return", f"{v.upside:+.1%}", GOOD if v.upside > 0 else BAD)]):
        y = 1.85 + i * 1.6
        d.rect(s, 0.6, y, 3.4, 1.4, PANEL)
        d.text(s, 0.85, y + 0.15, 3, 0.3, lab.upper(), size=11, color=MUTED, bold=True)
        d.text(s, 0.85, y + 0.5, 3, 0.75, val, size=34, color=col, bold=True, font=HEAD)
    thesis = av.get("thesis") or ins["thesis"]
    d.text(s, 4.5, 1.85, 8.2, 0.35, "Investment thesis" + ("" if av.get("thesis") else " (model-derived)"), size=16, color=NAVY, bold=True)
    d.bullets(s, 4.5, 2.3, 8.2, 3.2, thesis, size=14, gap=10)
    d.rect(s, 4.5, 5.55, 8.2, 1.2, PANEL)
    d.text(s, 4.7, 5.65, 7.9, 1.0, [[("Valuation: ", {"bold": True, "color": NAVY})] + [(" ".join(ins["valuation"][:2]), {})]], size=13)
    d.footer(s)

    # 3 ---------------- company snapshot
    s = d.slide()
    ann = f.annual.tail(5); rat = b.ratios[t].tail(5)
    g_last = rat["revenue_growth"].iloc[-1]
    d.title(s, f"{f.name.split(',')[0]}: {_bn(ann['revenue'].iloc[-1])} revenue, {g_last:+.1%} growth and a "
               f"{rat['ebit_margin'].iloc[-1]:.1%} EBIT margin in FY{int(ann.index[-1])}", "Company snapshot")
    rows = [["$m"] + [f"FY{int(y)}" for y in ann.index]]
    for lab, vals in [("Revenue", ann["revenue"] / 1e6), ("Growth", rat["revenue_growth"]), ("Gross margin", rat["gross_margin"]),
                      ("EBIT", ann["operating_income"] / 1e6), ("EBIT margin", rat["ebit_margin"]),
                      ("Net income", ann["net_income"] / 1e6), ("Free cash flow", ann["fcf"] / 1e6), ("ROIC", rat["roic"]),
                      ("Net debt / EBITDA", rat["net_debt_ebitda"])]:
        pct = lab in ("Growth", "Gross margin", "EBIT margin", "ROIC")
        rows.append([lab] + [("n/a" if np.isnan(x) else (_pct(x) if pct else (f"{x:.1f}x" if "/" in lab else f"{x:,.0f}"))) for x in vals])
    d.table(s, 0.6, 1.85, 7.9, rows, col_w=[2.15] + [1.15] * 5, size=12, row_h=0.42)
    tc = v.comps.loc[t]
    kv = [("Market cap", _bn(v.market_cap)), ("Enterprise value", _bn(v.market_cap + v.net_debt)), ("Net debt / (cash)", _bn(v.net_debt)),
          ("EV / EBITDA", fmt(tc["ev_ebitda"], "x")), ("P / E", fmt(tc["pe"], "x")),
          ("52-week range", f"${v.week52[0]:,.2f} – ${v.week52[1]:,.2f}"), ("Fiscal year end", f"{f.fiscal_year_end[:2]}/{f.fiscal_year_end[2:]}"),
          ("Industry (SIC)", f"{f.sic} {f.sic_description.title()}")]
    d.rect(s, 8.9, 1.85, 3.85, 4.75, PANEL)
    for i, (k, val) in enumerate(kv):
        d.text(s, 9.1, 2.0 + i * 0.57, 1.7, 0.3, k, size=12, color=MUTED)
        d.text(s, 10.75, 2.0 + i * 0.57, 1.85, 0.3, val, size=13, color=INK, bold=True, align=PP_ALIGN.RIGHT)
    d.footer(s)

    # 4 ---------------- performance vs peers
    s = d.slide()
    gt, mt = b.trends["revenue_growth"], b.trends["ebit_margin"]
    m_gap = sc.loc["ebit_margin", "target"] - sc.loc["ebit_margin", "peer_median"]
    d.title(s, f"EBIT margin is {abs(m_gap) * 100:.1f}pp {'above' if m_gap > 0 else 'below'} the peer median; "
               f"3-year revenue CAGR {sc.loc['revenue_cagr_3y', 'target']:+.1%} vs {sc.loc['revenue_cagr_3y', 'peer_median']:+.1%} for peers",
            "Financial performance vs peers")
    cats = [f"FY{int(y)}" for y in gt.index]
    d.chart(s, XL_CHART_TYPE.LINE_MARKERS, 0.6, 1.8, 6.0, 4.4, cats, [(t, gt[t]), ("Peer median", gt["peer_median"])], [ACCENT, PEER],
            title="Revenue growth")
    cats = [f"FY{int(y)}" for y in mt.index]
    d.chart(s, XL_CHART_TYPE.LINE_MARKERS, 6.75, 1.8, 6.0, 4.4, cats, [(t, mt[t]), ("Peer median", mt["peer_median"])], [ACCENT, PEER],
            title="EBIT margin")
    d.text(s, 0.6, 6.35, 12.1, 0.5, ins["financial"][1] if len(ins["financial"]) > 1 else "", size=13, color=INK)
    d.footer(s)

    # 5 ---------------- scorecard
    s = d.slide()
    better = int((sc["verdict"] == "Better").sum()); worse = int((sc["verdict"] == "Worse").sum())
    sc_head = (f"{t} trails the peer median on {worse} of {len(sc)} metrics and leads on {better or 'none'}" if worse >= better
               else f"{t} beats the peer median on {better} of {len(sc)} metrics and trails on {worse or 'none'}")
    d.title(s, f"Benchmark scorecard: {sc_head}",
            "Peer benchmarking")
    rows = [["Metric", t, "Peer median", "Peer range (Q1–Q3)", "Percentile", "Verdict"]]
    fills = {}
    for i, (m, r) in enumerate(sc.iterrows()):
        k = r["format"]
        rows.append([r["label"], fmt(r["target"], k), fmt(r["peer_median"], k), f"{fmt(r['peer_q1'], k)} – {fmt(r['peer_q3'], k)}",
                     "n/a" if np.isnan(r["percentile"]) else f"{r['percentile']:.0f}", r["verdict"]])
        fills[(i + 1, 5)] = {"Better": "D7EFE2", "Worse": "F8D9D5", "In line": "EDEFF2"}.get(r["verdict"])
    d.table(s, 0.6, 1.8, 12.1, rows, col_w=[3.9, 1.6, 1.6, 2.5, 1.2, 1.3], size=12, row_h=0.4, fills=fills,
            aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * 4 + [PP_ALIGN.CENTER])
    d.footer(s, "Percentile: rank within the peer group, oriented so 100 = best (e.g. lower leverage scores higher). Latest fiscal year.")

    # 6 ---------------- industry trends
    s = d.slide()
    ind = b.industry
    trend = "expanding" if ind["peer_margin_change_3y"] > 0.005 else ("contracting" if ind["peer_margin_change_3y"] < -0.005 else "flat")
    d.title(s, f"Industry backdrop: peer revenue growth averaged {ind['peer_growth_avg_3y']:+.1%} a year and margins are {trend}",
            "Industry trends")
    gt = b.trends["revenue_growth"]; cats = [f"FY{int(y)}" for y in gt.index]
    d.chart(s, XL_CHART_TYPE.LINE_MARKERS, 0.6, 1.8, 7.2, 4.6, cats,
            [("Peer 75th percentile", gt["peer_q3"]), ("Peer median", gt["peer_median"]), ("Peer 25th percentile", gt["peer_q1"]), (t, gt[t])],
            ["A9B8CC", PEER, "A9B8CC", ACCENT], title="Revenue growth: peer distribution vs " + t)
    d.text(s, 8.2, 1.85, 4.5, 0.35, "What the peer group tells us", size=16, color=NAVY, bold=True)
    d.bullets(s, 8.2, 2.35, 4.5, 4.2, ins["industry"], size=14, gap=12)
    d.footer(s)

    # 7 ---------------- returns & cash
    s = d.slide()
    lat = b.latest.copy()
    order = lat["roic"].sort_values(ascending=False).index.tolist()
    roic_rank = order.index(t) + 1
    d.title(s, f"Returns and cash: {t} ranks #{roic_rank} of {len(order)} on ROIC and converts "
               f"{sc.loc['fcf_conversion', 'target']:.0%} of earnings into free cash flow", "Returns & cash generation")
    for k, (m, ttl, x) in enumerate([("roic", "ROIC (latest fiscal year)", 0.6), ("fcf_margin", "FCF margin (latest fiscal year)", 6.75)]):
        o = lat[m].sort_values(ascending=False)
        ch = d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, x, 1.8, 6.0, 4.7, list(o.index), [(ttl, o.values)], [PEER],
                     legend=False, labels=True, title=ttl, label_fmt="0%")
        for i, tk in enumerate(o.index):
            if tk == t:
                pt = ch.plots[0].series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(ACCENT)
    d.footer(s, f"{t} highlighted. ROIC = EBIT × (1 − tax) / average (debt + equity − cash). Source: SEC EDGAR; Ticker-to-Thesis analysis.")

    # 8 ---------------- DCF
    s = d.slide()
    dt = v.dcf.table
    d.title(s, f"DCF: ${v.dcf.price_perpetuity:,.2f} per share at a {a.wacc:.1%} WACC and {a.terminal_growth:.1%} terminal growth",
            "Valuation · discounted cash flow")
    cats = [f"Y{i}" for i in dt.index]
    d.chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.6, 1.8, 6.4, 3.4, cats, [("Unlevered FCF ($m)", dt["ufcf"] / 1e6)], [PEER],
            number_format='#,##0', legend=False, labels=True, title="Unlevered free cash flow ($m)", label_fmt='#,##0')
    rows = [["Forecast driver"] + cats,
            ["Revenue growth"] + [_pct(x) for x in dt["growth"]], ["EBIT margin"] + [_pct(x) for x in dt["ebit_margin"]],
            ["Revenue ($m)"] + [f"{x / 1e6:,.0f}" for x in dt["revenue"]]]
    d.table(s, 0.6, 5.35, 6.4, rows, col_w=[1.9] + [0.9] * 5, size=11, row_h=0.36)
    br = [["EV bridge ($m)", "Value"], ["PV of forecast UFCF", f"{v.dcf.pv_fcf / 1e6:,.0f}"],
          ["PV of terminal value", f"{v.dcf.pv_tv_perpetuity / 1e6:,.0f}"], ["Enterprise value", f"{v.dcf.ev_perpetuity / 1e6:,.0f}"],
          ["Less: net debt", f"{v.net_debt / 1e6:,.0f}"], ["Equity value", f"{(v.dcf.ev_perpetuity - v.net_debt) / 1e6:,.0f}"],
          ["Diluted shares (m)", f"{v.shares / 1e6:,.1f}"], ["Value per share", _money(v.dcf.price_perpetuity)],
          ["Terminal value % of EV", _pct(v.dcf.tv_share)], ["Implied exit EV/EBITDA", f"{v.dcf.implied_exit_multiple:.1f}x"],
          ["Exit-multiple method", f"{_money(v.dcf.price_exit)} at {a.exit_multiple:.1f}x"]]
    d.table(s, 7.4, 1.8, 5.35, br, col_w=[3.15, 2.2], size=12, row_h=0.42, highlight_rows=(3, 7))
    d.footer(s, "UFCF = EBIT × (1 − tax) + D&A − capex − Δ working capital; mid-year discounting. Full model in the Excel workbook.")

    # 9 ---------------- WACC
    s = d.slide(); w = v.wacc
    d.title(s, f"Cost of capital: {w.wacc:.1%} WACC from a bottom-up beta of {w.beta_relevered:.2f} and an implied {w.rating} credit rating",
            "Valuation · WACC")
    steps = [("Risk-free rate", _pct(w.rf), w.rf_source), ("Equity risk premium", _pct(w.erp), "Assumption (config)"),
             ("Unlevered beta (peer median)", f"{w.beta_unlevered_median:.2f}", f"{len(w.beta_table) - 1} peers, 2y weekly vs S&P 500, Blume-adjusted"),
             ("Relevered beta", f"{w.beta_relevered:.2f}", f"At {t}'s market D/E; own beta {w.beta_own_adjusted:.2f}"),
             ("Cost of equity", _pct(w.cost_of_equity), "CAPM: rf + β × ERP"),
             ("Cost of debt (pre-tax)", _pct(w.cost_of_debt_pre), f"rf + {w.spread:.2%} spread ({w.rating}; coverage {w.coverage:.1f}x)"),
             ("Weights (equity / debt)", f"{w.weight_equity:.0%} / {w.weight_debt:.0%}", "Market cap and book debt"),
             ("WACC", _pct(w.wacc), "")]
    rows = [["Component", "Value", "Basis"]] + [[a_, b_, c_] for a_, b_, c_ in steps]
    d.table(s, 0.6, 1.8, 7.6, rows, col_w=[2.8, 1.3, 3.5], size=12, row_h=0.46, highlight_rows=(8,),
            aligns=[PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.LEFT])
    bt = w.beta_table.sort_values("unlevered_beta", ascending=False)
    ch = d.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, 8.6, 1.8, 4.15, 4.8, list(bt.index), [("Unlevered beta", bt["unlevered_beta"])], [PEER],
                 number_format='0.00', legend=False, labels=True, title="Unlevered betas", label_fmt='0.00', gap=50)
    ch.category_axis.reverse_order = True
    for i, tk in enumerate(bt.index):
        if tk == t:
            p_ = ch.plots[0].series[0].points[i]; p_.format.fill.solid(); p_.format.fill.fore_color.rgb = rgb(ACCENT)
    d.footer(s)

    # 10 ---------------- sensitivity
    s = d.slide()
    sg = v.sens_growth
    d.title(s, f"Sensitivity: every ±0.5pp on WACC moves the DCF value by about ${abs(sg.iloc[1, 2] - sg.iloc[3, 2]) / 2:,.0f} per share",
            "Valuation · sensitivity")
    rows = [["WACC \\ terminal growth"] + [f"{g:.2%}" for g in sg.columns]]
    fills = {}
    for i, (wv, row) in enumerate(sg.iterrows()):
        rows.append([f"{wv:.2%}"] + [_money(x) for x in row.values])
        for j, x in enumerate(row.values):
            fills[(i + 1, j + 1)] = "D7EFE2" if x >= v.price * 1.15 else ("F8D9D5" if x <= v.price * 0.9 else "FFFFFF")
    fills[(3, 3)] = "FFE2C2"
    d.table(s, 0.6, 1.9, 7.2, rows, col_w=[2.2] + [1.0] * 5, size=13, row_h=0.55, fills=fills,
            aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 5)
    d.text(s, 0.6, 5.35, 7.2, 0.6, f"Green: ≥15% above the current price (${v.price:,.2f}). Red: ≥10% below. Orange: base case.", size=12, color=MUTED)
    sx = v.sens_exit
    rows = [["WACC \\ exit multiple"] + [f"{m:.1f}x" for m in sx.columns]]
    for wv, row in sx.iterrows(): rows.append([f"{wv:.2%}"] + [_money(x) for x in row.values])
    d.text(s, 8.2, 1.9, 4.5, 0.3, "Exit-multiple method", size=14, color=NAVY, bold=True)
    d.table(s, 8.2, 2.3, 4.55, rows, col_w=[1.15] + [0.68] * 5, size=11, row_h=0.45, aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 5,
            fills={(3, 3): "FFE2C2"})
    d.footer(s)

    # 11 ---------------- comps
    s = d.slide()
    cp = v.comps; pm = cp.loc[~cp["is_target"], "ev_ebitda"].median(); te = cp.loc[t, "ev_ebitda"]
    prem = te / pm - 1 if pm and not np.isnan(te) else np.nan
    d.title(s, f"Trading comps: {t} trades at {te:.1f}x EV/EBITDA, a {abs(prem):.0%} {'premium' if prem > 0 else 'discount'} to the peer median of {pm:.1f}x",
            "Valuation · trading comparables")
    rows = [["Company", "Mkt cap", "EV", "EV / Sales", "EV / EBITDA", "P / E", "EBIT margin", "FCF yield"]]
    order = [x for x in cp.sort_values("market_cap", ascending=False).index if x != t] + [t]
    for tk in order:
        r = cp.loc[tk]
        rows.append([f"{tk} · {r['name'][:22]}", _bn(r["market_cap"]), _bn(r["enterprise_value"]), fmt(r["ev_revenue"], "x"),
                     fmt(r["ev_ebitda"], "x"), fmt(r["pe"], "x"), _pct(r["ebit_margin"]), _pct(r["fcf_yield"])])
    pe = cp[~cp["is_target"]]
    rows.append(["Peer median", "", "", fmt(pe["ev_revenue"].median(), "x"), fmt(pe["ev_ebitda"].median(), "x"), fmt(pe["pe"].median(), "x"),
                 _pct(pe["ebit_margin"].median()), _pct(pe["fcf_yield"].median())])
    d.table(s, 0.6, 1.8, 12.1, rows, col_w=[3.6, 1.15, 1.15, 1.2, 1.3, 1.1, 1.3, 1.3], size=12, row_h=0.42,
            highlight_rows=(len(rows) - 2,), fills={(len(rows) - 1, j): "E3E8F0" for j in range(8)})
    imp = v.comps_implied
    msg = " · ".join(f"{lab}: {_money(imp[k]['median'])} (range {_money(imp[k]['low'])}–{_money(imp[k]['high'])})"
                     for k, lab in [("ev_ebitda", "EV/EBITDA"), ("pe", "P/E")] if k in imp)
    d.text(s, 0.6, 1.8 + 0.42 * len(rows) + 0.25, 12.1, 0.6, [[("Implied value per share for " + t + ": ", {"bold": True, "color": NAVY}), (msg, {})]], size=13)
    d.footer(s, "Multiples on latest reported fiscal year; EV uses the latest balance sheet. Negative or extreme multiples excluded.")

    # 12 ---------------- football field (drawn with shapes for exact reference lines)
    s = d.slide()
    ff = v.football
    d.title(s, f"Valuation range: methods point to ${ff['low'].min():,.0f}–${ff['high'].max():,.0f}; price target ${v.price_target:,.2f}",
            "Valuation · football field")
    lo = min(ff["low"].min(), v.price, v.price_target) * 0.9; hi = max(ff["high"].max(), v.price, v.price_target) * 1.05
    x0, x1, y0, rowh = 4.6, 12.4, 2.25, 0.62
    X = lambda val: x0 + (val - lo) / (hi - lo) * (x1 - x0)
    for k in range(6):
        tick = lo + k * (hi - lo) / 5
        d.text(s, X(tick) - 0.5, y0 + rowh * len(ff) + 0.05, 1.0, 0.3, f"${tick:,.0f}", size=11, color=MUTED, align=PP_ALIGN.CENTER)
    for i, r in ff.iterrows():
        y = y0 + i * rowh
        d.text(s, 0.6, y + 0.12, 3.9, 0.4, r["method"], size=12, color=INK)
        d.rect(s, X(r["low"]), y + 0.12, max(X(r["high"]) - X(r["low"]), 0.04), 0.38, PEER if i else NEUTRAL, name=f"Range {i}")
        d.text(s, X(r["low"]) - 0.95, y + 0.17, 0.9, 0.3, f"${r['low']:,.0f}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)
        d.text(s, X(r["high"]) + 0.05, y + 0.17, 0.9, 0.3, f"${r['high']:,.0f}", size=11, color=MUTED)
    ybot = y0 + rowh * len(ff)
    d.vline(s, X(v.price), y0 - 0.1, ybot, INK, 1.5, dash=True)
    d.vline(s, X(v.price_target), y0 - 0.1, ybot, ACCENT, 2.5)
    left_first = v.price <= v.price_target
    for val, lab, col, right in [(v.price, f"Price ${v.price:,.2f}", INK, not left_first), (v.price_target, f"Target ${v.price_target:,.2f}", ACCENT, left_first)]:
        d.text(s, X(val) + (0.08 if right else -2.08), y0 - 0.42, 2.0, 0.3, lab, size=12, color=col, bold=True,
               align=PP_ALIGN.LEFT if right else PP_ALIGN.RIGHT)
    comp = v.pt_components
    d.text(s, 0.6, 6.35, 12.1, 0.5, "Price target = " + " + ".join(f"{wt:.0%} × {lab} ({_money(val)})" for lab, (val, wt) in
           zip(["DCF perpetuity", "DCF exit multiple", "comps EV/EBITDA"], comp.values()) if wt > 0), size=12, color=MUTED)
    d.footer(s)

    # 13 ---------------- scenarios
    s = d.slide()
    sc_ = v.scenarios; cfgs = rep.cfg["scenarios"]
    d.title(s, f"Scenarios: ${sc_['Bear'].price_perpetuity:,.0f} bear, ${sc_['Base'].price_perpetuity:,.0f} base, "
               f"${sc_['Bull'].price_perpetuity:,.0f} bull (DCF value per share)", "Scenario analysis")
    for i, (name, col, shift) in enumerate([("Bear", BAD, cfgs["bear"]), ("Base", NAVY, {"growth": 0, "margin": 0, "wacc": 0, "terminal_growth": 0}),
                                            ("Bull", GOOD, cfgs["bull"])]):
        x = 0.6 + i * 4.1; r = sc_[name]
        d.rect(s, x, 1.85, 3.8, 4.7, PANEL)
        d.text(s, x + 0.3, 2.0, 3.2, 0.4, name.upper() + " CASE", size=13, color=col, bold=True)
        d.text(s, x + 0.3, 2.45, 3.2, 0.8, _money(r.price_perpetuity), size=36, color=col, bold=True, font=HEAD)
        d.text(s, x + 0.3, 3.3, 3.2, 0.35, f"{r.price_perpetuity / v.price - 1:+.0%} vs current price", size=13, color=MUTED)
        items = [("Revenue growth", f"{a.revenue_growth[0] + shift['growth']:+.1%} → {a.revenue_growth[-1] + shift['growth']:+.1%}"),
                 ("EBIT margin (Y5)", _pct(a.ebit_margin[-1] + shift["margin"])), ("WACC", _pct(a.wacc + shift["wacc"])),
                 ("Terminal growth", _pct(a.terminal_growth + shift["terminal_growth"]))]
        for k, (lab, val) in enumerate(items):
            d.text(s, x + 0.3, 3.95 + k * 0.6, 1.9, 0.3, lab, size=13, color=MUTED)
            d.text(s, x + 2.0, 3.95 + k * 0.6, 1.5, 0.3, val, size=13, color=INK, bold=True, align=PP_ALIGN.RIGHT)
    d.footer(s)

    # 14 ---------------- risks
    s = d.slide()
    d.title(s, "Key risks and what would change our view", "Risks")
    auto = ins["risks"] or ["No quantitative red flags triggered by the screening rules."]
    d.text(s, 0.6, 1.85, 5.9, 0.35, "Flagged by the data", size=16, color=NAVY, bold=True)
    d.bullets(s, 0.6, 2.35, 5.9, 4.4, auto[:5], size=13, gap=10, marker_color=BAD)
    right = av.get("risks") or ins["weaknesses"] or ["No metric in the bottom 40% of the peer group."]
    d.text(s, 6.9, 1.85, 5.8, 0.35, "Analyst view" if av.get("risks") else "Where the company trails its peers", size=16, color=NAVY, bold=True)
    d.bullets(s, 6.9, 2.35, 5.8, 2.6, right[:4], size=13, gap=10, marker_color=BAD)
    cat = av.get("catalysts") or ins["strengths"]
    if cat:
        d.text(s, 6.9, 4.85, 5.8, 0.35, "Catalysts" if av.get("catalysts") else "Where the company leads", size=16, color=NAVY, bold=True)
        d.bullets(s, 6.9, 5.3, 5.8, 1.5, cat[:3], size=13, gap=8, marker_color=GOOD)
    d.footer(s)

    # 15 ---------------- methodology
    s = d.slide()
    d.title(s, "Methodology, data and limitations", "Appendix")
    meth = [("Data.", f"10-K/10-Q XBRL facts from SEC EDGAR for {t} and {len(peers)} peers, standardised across tag changes and restatements; "
                     "every value traceable to its filing (Excel: Sources sheet). Prices: Yahoo Finance / Stooq; risk-free: FRED DGS10."),
            ("Benchmarking.", "Latest fiscal year for each company; percentile ranks oriented so 100 = best; trends use the peer median and quartiles."),
            ("DCF.", f"{a.years}-year forecast, unlevered FCF, mid-year discounting, Gordon growth and exit-multiple terminal values."),
            ("WACC.", "CAPM with a bottom-up beta (peer regression betas, Blume-adjusted, unlevered and relevered); cost of debt from a synthetic rating."),
            ("Price target.", "Weighted blend of DCF (perpetuity and exit multiple) and trading comps; rating thresholds +15% / −10%."),
            ("Limitations.", "No consensus estimates or stub-period adjustment; multiples on the latest fiscal year, not LTM; operating leases "
                             "excluded from debt; peer set chosen by the analyst.")]
    d.bullets(s, 0.6, 1.85, 12.1, 4.6, meth, size=13, gap=9)
    d.text(s, 0.6, 6.45, 12.1, 0.4, "Educational research project generated with Ticker-to-Thesis. Not investment advice.", size=12, color=MUTED)
    d.footer(s)

    path = Path(path); d.prs.save(path)
    return path

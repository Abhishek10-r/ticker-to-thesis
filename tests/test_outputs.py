import shutil, subprocess, sys
from pathlib import Path
import pytest
from openpyxl import load_workbook
from pptx import Presentation

RECALC = Path("/mnt/skills/public/xlsx/scripts/recalc.py")


def test_deck_builds(report, tmp_path):
    from t2t.deck import build_deck
    p = build_deck(report, tmp_path / "deck.pptx")
    prs = Presentation(p)
    assert len(prs.slides) == 15
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert f"${report.val.price_target:,.2f}" in text


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice needed to recalculate formulas")
def test_excel_formulas_match_python(report, tmp_path):
    from t2t.excel import build_workbook
    p = build_workbook(report, tmp_path / "model.xlsx")
    script = RECALC if RECALC.exists() else None
    if script is None: pytest.skip("recalc helper not available")
    out = subprocess.run([sys.executable, str(script), str(p), "120"], capture_output=True, text=True).stdout
    assert '"total_errors": 0' in out
    wb = load_workbook(p, data_only=True); v = report.val
    assert wb["DCF"]["C32"].value == pytest.approx(v.dcf.price_perpetuity, rel=1e-9)
    assert wb["DCF"]["C42"].value == pytest.approx(v.dcf.price_exit, rel=1e-9)
    assert wb["DCF"]["C23"].value == pytest.approx(v.wacc.wacc, rel=1e-9)
    assert wb["Valuation"]["C17"].value == pytest.approx(v.price_target, rel=1e-9)
    assert wb["Valuation"]["C20"].value == v.rating
    assert wb["Sensitivity"]["C8"].value == pytest.approx(v.sens_growth.iloc[0, 0], rel=1e-9)
